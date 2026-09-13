"""Rebuild the Word project manual from the maintained Markdown specifications."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.7)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
for name, size in [('Title', 28), ('Heading 1', 20), ('Heading 2', 15), ('Heading 3', 12)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].paragraph_format.keep_with_next = True

def inline(p, text):
    text = re.sub(r'\[([^]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        r = p.add_run(part.strip('*`') if part.startswith(('**', '`')) else part)
        if part.startswith('**'):
            r.bold = True
        if part.startswith('`'):
            r.font.name = 'Consolas'
            r.font.size = Pt(9)

doc.add_paragraph('Packers Nepal Project Manual', 'Title')
doc.add_paragraph('Business requirements and implementation plan', 'Subtitle')
doc.add_paragraph('Prepared for the business owner and development team\n11 September 2026')
doc.add_paragraph('This manual defines a packing services application using Laravel 12, Filament, Livewire, and Tailwind CSS. It explains the service boundaries, customer and staff processes, page-level access rules, database design, and the work required to deliver and operate the application.')
doc.add_paragraph('The business prepares and protects goods for customers and their chosen transporters. Transportation is outside the proposed baseline service. The chapters describe the target system; implementation progress is tracked separately in the repository README.')
doc.add_heading('Contents', 1)
for line in ['1  Business specification and process', '2  Roles permissions and page assignments', '3  Architecture and implementation roadmap']:
    doc.add_paragraph(line)

for filename in ['PROJECT-SPECIFICATION.md', 'ACCESS-CONTROL.md', 'IMPLEMENTATION.md']:
    doc.add_page_break()
    lines = (ROOT / filename).read_text(encoding='utf-8').splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith('```'):
            lang = line[3:]
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                block.append(lines[i]); i += 1
            if lang == 'mermaid':
                for step in ['Customer requests a quote', 'Sales qualifies the request and performs any assessment', 'Staff prepare and send an itemized quote', 'Customer accepts or requests a revision', 'Operations confirms crew capacity and materials', 'Assigned team packs labels and checks the goods', 'Customer or nominated contact acknowledges handover', 'Finance reconciles the invoice and closes the job']:
                    doc.add_paragraph(step, 'List Number')
            else:
                for code in block:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    r = p.add_run(code)
                    r.font.name, r.font.size = 'Consolas', Pt(8)
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                row = [x.strip() for x in lines[i].strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+', x) for x in row): rows.append(row)
                i += 1
            i -= 1
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.autofit = False
            ratios = {2: [.31,.69], 3: [.23,.49,.28], 4: [.18,.34,.34,.14]}.get(len(rows[0]), [1/len(rows[0])]*len(rows[0]))
            for col, ratio in zip(table.columns, ratios): col.width = Inches(6.87*ratio)
            for n, row in enumerate(rows):
                cells = table.add_row().cells
                trpr = table.rows[-1]._tr.get_or_add_trPr()
                trpr.append(OxmlElement('w:cantSplit'))
                if n == 0: trpr.append(OxmlElement('w:tblHeader'))
                for j, value in enumerate(row):
                    cell = cells[j]
                    cell.width = Inches(6.87*ratios[j])
                    cell.vertical_alignment = 1
                    pr = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement('w:tcBorders')
                    for edge in ['top','left','bottom','right']:
                        el = OxmlElement('w:'+edge)
                        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]: el.set(qn('w:'+k),v)
                        borders.append(el)
                    pr.append(borders)
                    sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'DDE8ED' if n==0 else ('F6F8F9' if n%2==0 else 'FFFFFF')); pr.append(sh)
                    margins = OxmlElement('w:tcMar')
                    for edge in ['top','left','bottom','right']:
                        el=OxmlElement('w:'+edge); el.set(qn('w:w'),'90'); el.set(qn('w:type'),'dxa'); margins.append(el)
                    pr.append(margins)
                    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(3)
                    inline(p,value)
                    for r in p.runs: r.font.size=Pt(9); r.bold = n==0 or r.bold
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif line.startswith('#'):
            level = len(line)-len(line.lstrip('#'))
            title = re.sub(r'[^\w\s]', '', line.lstrip('# ').strip())
            doc.add_heading(title, min(level,3))
        elif line.strip():
            style = 'Normal'
            if line.startswith('- '): line=line[2:]; style='List Bullet'
            elif re.match(r'^\d+\. ',line): line=re.sub(r'^\d+\. ','',line); style='List Number'
            inline(doc.add_paragraph(style=style), line)
        i += 1

footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('Packers Nepal  |  ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.core_properties.title='Packers Nepal Project Manual'
doc.core_properties.subject='Packing services requirements and implementation'
doc.core_properties.author='Packers Nepal'
output=ROOT/'Packers-Nepal-Project-Manual.docx'
doc.save(output)
print(output)
